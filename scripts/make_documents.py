"""Create editable Word content sources and templates without replacing existing files."""
import json
import re
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent


def document(title):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(20)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption', 'List Bullet', 'List Number']:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.15
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Title'].paragraph_format.space_after = Pt(16)
    doc.styles['Heading 1'].font.size = Pt(13)
    doc.styles['Heading 1'].font.bold = True
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(14)
    doc.styles['Heading 1'].paragraph_format.keep_with_next = True
    doc.add_paragraph(title, 'Title')
    doc.core_properties.author = 'Mohammad Mahmudul Hasan'
    return doc


def link(paragraph, label, target):
    relationship = paragraph.part.relate_to(target, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    element = OxmlElement('w:hyperlink'); element.set(qn('r:id'), relationship)
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '24649F'); props.append(color)
    run.append(props)
    text = OxmlElement('w:t'); text.text = label; run.append(text)
    element.append(run); paragraph._p.append(element)


def rich(paragraph, value):
    tokens = re.split(r'(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|\*[^*]+\*)', value)
    for token in tokens:
        if not token:
            continue
        match = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
        if match:
            link(paragraph, match.group(1), match.group(2))
        elif token.startswith('**') and token.endswith('**'):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith('*') and token.endswith('*'):
            paragraph.add_run(token[1:-1]).italic = True
        else:
            paragraph.add_run(token)


def add_blocks(doc, blocks):
    for block in blocks:
        if block['type'] == 'text':
            if block['heading']:
                doc.add_paragraph(block['heading'], 'Heading 1')
            for value in re.split(r'\n\s*\n', block['body']):
                if not value.strip():
                    continue
                for line in value.splitlines():
                    style = 'Normal'
                    if line.startswith('- '):
                        style, line = 'List Bullet', line[2:]
                    paragraph = doc.add_paragraph(style=style)
                    rich(paragraph, line)
        elif block['type'] == 'image':
            picture = doc.add_paragraph().add_run().add_picture(str(ROOT / block['src'].lstrip('/')), width=Mm(155 if block.get('size') == 'full' else 110))
            picture._inline.docPr.set('descr', block['alt'])
            if block['caption']:
                doc.add_paragraph('Caption: ' + block['caption'], 'Caption')
            doc.add_paragraph('Image width: ' + block.get('size', 'full'))


def save(doc, path):
    if path.exists():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    print(path.relative_to(ROOT))


def create():
    projects = json.loads((ROOT / '_data/projects.json').read_text(encoding='utf-8'))
    research = json.loads((ROOT / '_data/research.json').read_text(encoding='utf-8'))
    for key, item in projects.items():
        doc = document(item['title'])
        for label, value in [('Title', item['title']), ('Summary', item['summary']), ('Context', item['context']), ('Role', item['role']), ('Order', item['order']), ('Featured', 'yes' if item['featured'] else 'no'), ('Related research', item['research_anchor'])]:
            doc.add_paragraph(f'{label}: {value}')
        add_blocks(doc, item['blocks'])
        save(doc, ROOT / 'content/projects' / key / 'project.docx')
    for item in research['modules']:
        doc = document(item['title'])
        doc.add_paragraph('Title: ' + item['title'])
        doc.add_paragraph('Summary: ' + item['summary'])
        add_blocks(doc, item['blocks'])
        save(doc, ROOT / 'content/research' / item['id'] / 'research.docx')
    doc = document('Project content template')
    doc.add_paragraph('Note: Fill the labelled fields and replace every bracketed instruction. Keep section headings in the Word Heading 1 style.')
    for label, value in [('Title', '[REPLACE with your project title]'), ('Summary', '[REPLACE with a one or two sentence summary]'), ('Context', '[REPLACE with employer or academic context]'), ('Role', '[REPLACE with your individual contribution]'), ('Order', '8'), ('Featured', 'no'), ('Related research', '')]:
        doc.add_paragraph(f'{label}: {value}')
    for heading, body in [('Problem', '[REPLACE with the problem and why it matters]'), ('My contribution', '[REPLACE with what you personally developed]'), ('Method', '[REPLACE with the approach and tools used]'), ('Evaluation', '[REPLACE with results and evaluation context]'), ('Limitations and next steps', '[REPLACE with limitations and a future research question]')]:
        doc.add_paragraph(heading, 'Heading 1'); doc.add_paragraph(body)
    save(doc, ROOT / 'content/templates/project-template.docx')
    doc = document('Research content template')
    doc.add_paragraph('Note: Replace the bracketed instructions. Use Heading 1 for sections. Insert pictures In Line with Text; put Caption: and Alt: paragraphs immediately below each picture.')
    doc.add_paragraph('Title: [REPLACE with the research direction]')
    doc.add_paragraph('Summary: [REPLACE with a short introduction]')
    for heading, body in [('Relevant experience', '[REPLACE with completed work that supports this direction]'), ('Research question', '[REPLACE with the question you want to investigate]'), ('Proposed approach', '[REPLACE with the proposed methods and how you would evaluate them]')]:
        doc.add_paragraph(heading, 'Heading 1'); doc.add_paragraph(body)
    save(doc, ROOT / 'content/templates/research-template.docx')


if __name__ == '__main__':
    create()
